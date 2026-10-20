from pptx import Presentation
from pptx.util import Emu

prs = Presentation(r"c:\proyectos\bluebird\ke-simulator\docs\BBPPTX.pptx")
print("slide_size:", Emu(prs.slide_width).inches, "x", Emu(prs.slide_height).inches)
print("num_slides:", len(prs.slides))
print("num_layouts:", len(prs.slide_layouts))
for i, layout in enumerate(prs.slide_layouts):
    names = [ph.placeholder_format.idx for ph in layout.placeholders]
    phs = [(ph.placeholder_format.idx, ph.placeholder_format.type, ph.name) for ph in layout.placeholders]
    print(f"layout[{i}] name={layout.name!r} placeholders={phs}")
print("--- existing slides ---")
for i, slide in enumerate(prs.slides):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip()[:60])
    print(f"slide[{i}] layout={slide.slide_layout.name!r} texts={texts}")
